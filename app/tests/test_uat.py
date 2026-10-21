"""UAT: tile acceptance, page sign-off (never blocked; reason required with gaps), staleness, storage, export."""
import io
import sys
import unittest
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parent / "streamlit"))
sys.path.insert(0, str(HERE))

import uat  # noqa: E402
import uat_export  # noqa: E402
from fake_snowflake import FakeSession  # noqa: E402
from store import Store  # noqa: E402
from test_composer import APP  # noqa: E402

SPEC = {"pages": [{"page_id": "p1", "label": "Brief", "purpose": "Are we on plan?", "components": [
    {"component_id": "t1", "type": "kpi_plan", "title": "YTD vs plan", "question": "Are we on plan?"},
    {"component_id": "t2", "type": "driver_bridge", "title": "Drivers", "question": "What moved?", "metric_id": "M"},
    {"component_id": "t3", "type": "action_table", "title": "Decisions"}]}],
    "requirements": {"p1/t1": "BRD-4.2"}}
TILES = {"p1/t1": {"status": "pass", "status_eff": "pass", "sig": "a1", "title": "YTD vs plan",
                   "checks": [{"id": "QA-01", "status": "pass", "summary": "ok",
                               "rows": [{"label": "Total", "column": "Sales", "dashboard": 10, "independent": 10, "ok": True}],
                               "sql": {"dashboard": "SELECT 1", "independent": "SELECT 1"}}]},
         "p1/t2": {"status": "mismatch", "status_eff": "mismatch", "sig": "b1", "title": "Drivers",
                   "checks": [{"id": "QA-01", "status": "mismatch", "summary": "2 differ",
                               "rows": [{"label": "East", "column": "Sales", "dashboard": 10, "independent": 11, "ok": False}],
                               "sql": {"dashboard": "SELECT 2", "independent": "SELECT 2"}}]}}
DATA = (lambda pg, c: c["type"] != "action_table")


def tile_e(cid, decision, sig, comment="", at="2026-10-16 09:00"):
    return {"kind": "tile", "page_id": "p1", "tile_key": cid, "decision": decision, "tile_sig": sig,
            "comment": comment, "created_by": "SHARATH", "created_at": at}


class UatTests(unittest.TestCase):
    def test_not_started_and_non_data_tiles_skipped(self):
        st = uat.status(SPEC, TILES, [], DATA)
        self.assertEqual(st["status"], "not_started")
        pg = st["pages"][0]
        self.assertEqual([r["cid"] for r in pg["tiles"]], ["p1/t1", "p1/t2"])
        self.assertEqual(pg["tiles"][0]["requirement"], "BRD-4.2")
        self.assertEqual(pg["open_findings"], 1)

    def test_acceptance_reopens_when_numbers_change(self):
        st = uat.status(SPEC, TILES, [tile_e("p1/t1", "accepted", "a1")], DATA)
        self.assertEqual(st["pages"][0]["tiles"][0]["state"], "accepted")
        self.assertEqual(st["status"], "in_review")
        changed = {**TILES, "p1/t1": {**TILES["p1/t1"], "sig": "a2"}}
        self.assertEqual(uat.status(SPEC, changed, [tile_e("p1/t1", "accepted", "a1")], DATA)
                         ["pages"][0]["tiles"][0]["state"], "changed")

    def test_newest_entry_wins(self):
        es = [tile_e("p1/t1", "rejected", "a1", "wrong period", "2026-10-16 10:00"), tile_e("p1/t1", "accepted", "a1")]
        self.assertEqual(uat.status(SPEC, TILES, es, DATA)["pages"][0]["tiles"][0]["state"], "rejected")

    def test_tile_rules(self):
        with self.assertRaises(ValueError):
            uat.check_tile("rejected", "", "pass")                 # reject needs a reason
        with self.assertRaises(ValueError):
            uat.check_tile("accepted", "", "mismatch")             # accepting a mismatch needs a reason
        self.assertEqual(uat.check_tile("accepted", "known timing gap", "mismatch"), "accepted")
        self.assertEqual(uat.check_tile("accepted", "", "pass"), "accepted")

    def test_page_never_blocked_but_reason_required(self):
        pg = uat.status(SPEC, TILES, [], DATA)["pages"][0]          # 0 of 2 accepted, 1 open finding
        with self.assertRaises(ValueError):
            uat.check_page("approved", "", pg)
        self.assertEqual(uat.check_page("approved", "go-live; drivers gap owned by data team", pg), "approved_notes")
        with self.assertRaises(ValueError):
            uat.check_page("rejected", "", pg)
        clean = {**pg, "open_findings": 0, "accepted": 2}
        self.assertEqual(uat.check_page("approved", "", clean), "approved")

    def test_page_signoff_status_and_staleness(self):
        st = uat.status(SPEC, TILES, [tile_e("p1/t1", "accepted", "a1"), tile_e("p1/t2", "accepted", "b1", "known")], DATA)
        pg = st["pages"][0]
        self.assertEqual(pg["state"], "ready")
        sig = pg["sig"]
        es = [{"kind": "page", "page_id": "p1", "tile_key": "", "decision": "approved_notes", "page_sig": sig,
               "comment": "drivers gap owned", "created_by": "SHARATH", "created_at": "2026-10-16 11:00"},
              tile_e("p1/t1", "accepted", "a1"), tile_e("p1/t2", "accepted", "b1", "known")]
        st = uat.status(SPEC, TILES, es, DATA)
        self.assertEqual((st["status"], st["pages_signed"]), ("approved_notes", 1))
        changed = {**TILES, "p1/t2": {**TILES["p1/t2"], "sig": "b2"}}
        st2 = uat.status(SPEC, changed, es, DATA)
        self.assertEqual(st2["pages"][0]["state"], "changed")
        self.assertEqual(st2["status"], "in_review")
        es_rej = [{**es[0], "decision": "rejected"}] + es[1:]
        self.assertEqual(uat.status(SPEC, TILES, es_rej, DATA)["status"], "rejected")

    def test_store_round_trip(self):
        s = Store(FakeSession(), APP)
        s.actor = "Sharath"
        s.save_uat({"dashboard_id": "d1", "sv": "SV", "page_id": "p1", "tile_key": "p1/t1", "kind": "tile",
                    "decision": "accepted", "tile_sig": "a1"})
        s.save_uat({"dashboard_id": "d1", "sv": "SV", "page_id": "p1", "kind": "page", "decision": "approved",
                    "page_sig": "x", "comment": "ok"})
        rows = s.list_uat("SV", "d1")
        self.assertEqual([r["kind"] for r in rows], ["page", "tile"])          # newest first
        self.assertEqual(rows[1]["created_by"], "Sharath")
        self.assertEqual(s.list_uat("SV", "other"), [])

    def test_export_carries_acceptance_and_signoff(self):
        from openpyxl import load_workbook
        es = [{"kind": "page", "page_id": "p1", "tile_key": "", "decision": "approved_notes",
               "page_sig": uat.status(SPEC, TILES, [], DATA)["pages"][0]["sig"], "comment": "gap owned",
               "created_by": "SHARATH", "created_at": "2026-10-16 11:00"},
              tile_e("p1/t1", "accepted", "a1"), tile_e("p1/t2", "rejected", "b1", "wrong total")]
        st = uat.status(SPEC, TILES, es, DATA)
        import datetime as dt
        meta = {"title": "Brief – Dashboard UAT", "reference": "CHG1", "testers": "Sharath", "environment": "Snowflake",
                "semantic_view": "DB.S.SV", "base_tables": ["DB.S.T"], "period": "Nov 2024", "filters": "No filters",
                "description": "d", "recipients": "r", "started": dt.date(2026, 10, 16), "completed": dt.date(2026, 10, 16)}
        wb = load_workbook(io.BytesIO(uat_export.build(SPEC, {"tiles": TILES, "dashboard": []}, meta, uat_status=st)))
        self.assertEqual(wb.sheetnames, ["Test Cases", "Expected Results", "Requirement Trace", "Sign-off", "SQL Appendix"])
        meta["history"] = es
        wb2 = load_workbook(io.BytesIO(uat_export.build(SPEC, {"tiles": TILES, "dashboard": []}, meta, uat_status=st)))
        self.assertIn("Review History", wb2.sheetnames)
        hh = wb2["Review History"]
        acts = [hh.cell(row=r, column=4).value for r in range(2, hh.max_row + 1)]
        self.assertIn("Reviewer: needs change", acts)
        self.assertIn("Page signed off with notes", acts)
        tc = wb["Test Cases"]
        vals = [tc.cell(row=r, column=8).value for r in range(17, tc.max_row + 1)]
        self.assertIn("Fail", vals)                 # QA-01 mismatch and the rejected acceptance
        texts = " ".join(str(tc.cell(row=r, column=3).value) for r in range(17, tc.max_row + 1))
        self.assertIn("UAT business acceptance", texts)
        so = wb["Sign-off"]
        self.assertEqual((so["A2"].value, so["F2"].value, so["H2"].value), ("Brief", "Signed off with notes", "gap owned"))
        tr = wb["Requirement Trace"]
        self.assertEqual(tr["F2"].value, "BRD-4.2")
        self.assertTrue(str(tr["K3"].value).startswith("Needs change"))


class LoopTests(unittest.TestCase):
    def test_request_fix_recheck_cycle(self):
        es = [tile_e("p1/t1", "rejected", "a1", "show by region", "2026-10-12 09:00")]
        row = uat.status(SPEC, TILES, es, DATA)["pages"][0]["tiles"][0]
        self.assertEqual((row["state"], row["request"]["comment"]), ("rejected", "show by region"))
        changed = {**TILES, "p1/t1": {**TILES["p1/t1"], "sig": "a9"}}     # editing alone does not close it
        self.assertEqual(uat.status(SPEC, changed, es, DATA)["pages"][0]["tiles"][0]["state"], "rejected")
        es = [tile_e("p1/t1", "fixed", "a9", "now by region", "2026-10-13 09:00")] + es
        row = uat.status(SPEC, changed, es, DATA)["pages"][0]["tiles"][0]
        self.assertEqual(row["state"], "fixed")
        self.assertEqual([m["decision"] for m in row["thread"]], ["rejected", "fixed"])
        self.assertEqual(row["request"]["comment"], "show by region")
        es = [tile_e("p1/t1", "accepted", "a9", "", "2026-10-14 09:00")] + es
        row = uat.status(SPEC, changed, es, DATA)["pages"][0]["tiles"][0]
        self.assertEqual((row["state"], row["request"]), ("accepted", None))

    def test_builder_rules(self):
        with self.assertRaises(ValueError):
            uat.check_builder("fixed", "", "pending")                # nothing to fix
        with self.assertRaises(ValueError):
            uat.check_builder("replied", "", "rejected")             # a reply needs text
        self.assertEqual(uat.check_builder("fixed", "", "rejected"), "fixed")
        self.assertEqual(uat.check_builder("replied", "not in the view", "fixed"), "replied")

    def test_next_step_and_check_line(self):
        st = uat.status(SPEC, TILES, [], DATA)
        self.assertEqual(st["next"]["step"], 2)
        self.assertIn("review 2 tile(s)", st["next"]["text"].lower())
        ck = st["pages"][0]["tiles"][1]["check"]
        self.assertEqual(ck["text"], "Differs by 9.1%")
        self.assertIn("Dashboard 10.00 · Source 11.00", ck["detail"])
        none = uat.status(SPEC, {}, [], DATA)
        self.assertEqual(none["next"]["step"], 1)

    def test_inbox_across_dashboards(self):
        def e(did, cid, d, by, at, cm=""):
            return {"kind": "tile", "dashboard_id": did, "tile_key": cid, "page_id": cid.split("/")[0], "decision": d,
                    "created_by": by, "created_at": at, "comment": cm, "tile_sig": "s", "tile_title": cid,
                    "page_label": "Brief"}
        entries = [e("d2", "p1/t2", "fixed", "SHARATH", "2026-10-15 10:00", "done"),
                   e("d1", "p1/t1", "rejected", "PRIYA", "2026-10-12 09:00", "region view"),
                   e("d2", "p1/t2", "rejected", "PRIYA", "2026-10-14 09:00", "fiscal months"),
                   e("d3", "p1/t3", "accepted", "PRIYA", "2026-10-15 09:00"),
                   e("d3", "p1/t3", "rejected", "PRIYA", "2026-10-10 09:00", "x")]
        owners = {"d1": {"name": "Exec brief", "owner": "SHARATH"}, "d2": {"name": "Ops", "owner": "ANITA"},
                  "d3": {"name": "Pipe", "owner": "SHARATH"}}
        ib = uat.inbox(sorted(entries, key=lambda x: x["created_at"], reverse=True), owners, "SHARATH",
                       today=__import__("datetime").date(2026, 10, 16))
        self.assertEqual([(r["dashboard_name"], r["state"]) for r in ib["rows"]], [("Exec brief", "rejected"), ("Ops", "fixed")])
        self.assertEqual(ib["rows"][0]["age_days"], 4)
        self.assertEqual((ib["badge_requests"], ib["badge_recheck"], ib["closed_30d"]), (1, 0, 1))
        ib2 = uat.inbox(entries, owners, "PRIYA", today=__import__("datetime").date(2026, 10, 16))
        self.assertEqual(ib2["badge_recheck"], 1)

    def test_store_inbox_and_adopt(self):
        s = Store(FakeSession(), APP)
        s.actor = "PRIYA"
        s.save_uat({"dashboard_id": "", "sv": "SV", "page_id": "p1", "tile_key": "p1/t1", "kind": "tile",
                    "decision": "rejected", "comment": "x", "tile_title": "Sales"})
        s.save_uat({"dashboard_id": "d9", "sv": "SV", "page_id": "p1", "tile_key": "p1/t2", "kind": "tile",
                    "decision": "accepted"})
        s.adopt_uat("SV", "d1", "Exec brief")
        rows = s.list_uat_tiles()
        self.assertEqual([(r["dashboard_id"], r["tile_title"]) for r in rows], [("d1", "Sales")])
        s.set_setting("owner:d1", "ANITA")
        self.assertEqual(s.list_settings("owner:"), {"owner:d1": "ANITA"})


if __name__ == "__main__":
    unittest.main()
